from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from lab_meeting_slides.library import Library, build_template_deck, build_demo_deck

def test_components_and_template(tmp_path):
    p=tmp_path/'template.pptx'; build_template_deck(p)
    from pptx import Presentation
    prs=Presentation(p); assert len(prs.slides)==9
    demo=tmp_path/'demo.pptx'; build_demo_deck(demo); assert len(Presentation(demo).slides)==9

def test_widescreen():
    l=Library(); assert round(l.W/l.H,2)==1.78
