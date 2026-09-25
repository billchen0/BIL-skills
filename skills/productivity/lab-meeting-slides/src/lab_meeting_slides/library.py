"""Reusable, generic 16:9 lab-meeting slide components.

The library intentionally uses editable PowerPoint primitives rather than
flattened screenshots. Content, institutional marks, figures, and citations
belong to the consuming lab or project.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TOKENS = json.loads((ROOT / "tokens.json").read_text())


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.replace("#", "").upper())


def _bullet(paragraph, character: str = "•") -> None:
    props = paragraph._p.get_or_add_pPr()
    bullet = OxmlElement("a:buChar")
    bullet.set("char", character)
    props.insert(0, bullet)


def add_text(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 14,
    color: str = "text",
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.04,
    font_name: str = "Aptos",
):
    """Add a wrapped editable text box using inches for geometry."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(TOKENS["colors"].get(color, color))
    return shape


class Library:
    """Component and layout helpers for editable 16:9 PowerPoint slides."""

    WIDTH = 13.333
    HEIGHT = 7.5
    W = WIDTH
    H = HEIGHT

    def __init__(self, tokens: Mapping | None = None):
        self.tokens = tokens or TOKENS

    def color(self, token: str) -> RGBColor:
        value = self.tokens["colors"].get(token, token)
        return _rgb(value)

    def prs(self) -> Any:
        presentation = Presentation()
        presentation.slide_width = Inches(self.WIDTH)
        presentation.slide_height = Inches(self.HEIGHT)
        return presentation

    def background(self, slide, token: str = "canvas") -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.color(token)

    def rounded_box(
        self,
        slide,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: str = "surface",
        line: str = "line",
        radius: bool = True,
    ):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.color(fill)
        shape.line.color.rgb = self.color(line)
        shape.line.width = Pt(0.7)
        return shape

    def text(self, slide, value: str, *args, **kwargs):
        """Instance wrapper that honors a custom token dictionary."""
        shape = add_text(slide, value, *args, **kwargs)
        # add_text resolves global tokens; correct the run color for custom tokens.
        color = kwargs.get("color", "text")
        shape.text_frame.paragraphs[0].runs[0].font.color.rgb = self.color(color)
        return shape

    def pill(self, slide, label: str, x: float, y: float, w: float, h: float = 0.28, color: str = "teal"):
        self.rounded_box(slide, x, y, w, h, color, color)
        return self.text(
            slide,
            label.upper(),
            x,
            y + 0.035,
            w,
            h - 0.04,
            size=9,
            color="white",
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )

    def card(self, slide, x: float, y: float, w: float, h: float, label: str | None = None, accent: str = "blue"):
        shape = self.rounded_box(slide, x, y, w, h, "surface", "line")
        if label:
            self.pill(slide, label, x + 0.18, y + 0.16, min(1.7, w - 0.36), 0.24, accent)
        return shape

    def title_block(self, slide, title: str, subtitle: str | None = None, section: str | None = None):
        if section:
            self.pill(slide, section, 0.52, 0.38, 1.6, 0.28, "teal")
        self.text(slide, title, 0.52, 0.84, 10.6, 0.58, size=30, color="navy", bold=True)
        if subtitle:
            self.text(slide, subtitle, 0.55, 1.46, 11.0, 0.36, size=17, color="muted_text")

    def footer(self, slide, page: int, label: str = "LAB MEETING LIBRARY"):
        self.text(slide, label, 0.52, 7.10, 3.8, 0.2, size=9, color="muted_text", bold=True)
        self.rounded_box(slide, 12.42, 6.98, 0.42, 0.3, "blue", "blue")
        self.text(slide, f"{page:02d}", 12.42, 7.015, 0.42, 0.16, size=9, color="white", bold=True, align=PP_ALIGN.CENTER)

    def bullets(self, slide, items: Sequence[str], x: float, y: float, w: float, h: float, size: float = 14, color: str = "text"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.04)
        frame.margin_top = frame.margin_bottom = Inches(0.04)
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = item
            paragraph.level = 0
            paragraph.space_after = Pt(8)
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = self.color(color)
            _bullet(paragraph)
        return shape

    def callout(self, slide, label: str, body: str, x: float, y: float, w: float, h: float, fill: str = "soft_teal"):
        self.rounded_box(slide, x, y, w, h, fill, fill)
        self.text(slide, label.upper(), x + 0.18, y + 0.14, w - 0.36, 0.2, size=10, color="teal", bold=True)
        return self.text(slide, body, x + 0.18, y + 0.42, w - 0.36, h - 0.52, size=16, color="navy", bold=True)

    def metric(self, slide, value: str, label: str, x: float, y: float, w: float = 1.85, h: float = 1.0):
        self.card(slide, x, y, w, h)
        self.text(slide, value, x + 0.16, y + 0.17, w - 0.32, 0.36, size=26, color="navy", bold=True)
        return self.text(slide, label, x + 0.16, y + 0.62, w - 0.32, 0.2, size=10, color="muted_text")

    def figure_placeholder(self, slide, x: float, y: float, w: float, h: float, label: str = "FIGURE OR TABLE"):
        self.card(slide, x, y, w, h, "FIGURE", "blue")
        # Abstract, non-data-bearing visual rails.
        for index, height in enumerate((0.55, 0.9, 0.7, 1.15, 0.82)):
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x + 0.62 + index * 0.52),
                Inches(y + h - 0.72 - height),
                Inches(0.24),
                Inches(height),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = self.color("blue" if index % 2 == 0 else "teal")
            bar.line.fill.background()
        self.text(slide, label, x + 0.3, y + h / 2 - 0.15, w - 0.6, 0.3, size=11, color="muted_text", bold=True, align=PP_ALIGN.CENTER)

    def timeline(self, slide, items: Sequence[tuple[str, str]], x: float = 0.85, y: float = 3.2, w: float = 11.7):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = self.color("line")
        line.line.fill.background()
        step = w / max(1, len(items) - 1)
        for index, (heading, body) in enumerate(items):
            center = x + index * step
            node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center - 0.16), Inches(y - 0.16), Inches(0.32), Inches(0.32))
            node.fill.solid()
            node.fill.fore_color.rgb = self.color("canvas")
            node.line.color.rgb = self.color("teal" if index % 2 else "blue")
            node.line.width = Pt(4)
            self.text(slide, heading, center - 0.75, y + 0.42, 1.5, 0.28, size=12, color="navy", bold=True, align=PP_ALIGN.CENTER)
            self.text(slide, body, center - 0.88, y + 0.78, 1.76, 0.55, size=11, color="muted_text", align=PP_ALIGN.CENTER)

    def two_columns(self, slide, y: float = 2.05, h: float = 4.55):
        self.card(slide, 0.52, y, 6.02, h)
        self.card(slide, 6.79, y, 6.02, h)
        return (0.52, 6.79)

    def three_columns(self, slide, y: float = 2.05, h: float = 4.55):
        for x in (0.52, 4.51, 8.5):
            self.card(slide, x, y, 3.8, h)
        return (0.52, 4.51, 8.5)

    def add_picture(self, slide, path: str | Path, x: float, y: float, w: float, h: float):
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

    def add_chart(self, slide, chart_factory: Callable, x: float, y: float, w: float, h: float):
        """Insert a caller-owned native chart factory at the requested geometry."""
        return chart_factory(slide, x, y, w, h)

    def add_slide(self, presentation: Presentation, title: str, subtitle: str | None = None, section: str | None = None, page: int = 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.background(slide)
        self.title_block(slide, title, subtitle, section)
        self.footer(slide, page)
        return slide


def _build_deck(path: str | Path, demo: bool = False) -> Any:
    library = Library()
    presentation = library.prs()
    if demo:
        specs = [
            ("Study overview", "A generic sample deck for a wearable signal-quality pilot", "LAB MEETING"),
            ("Question and scope", "Why signal quality matters before downstream modeling", "SECTION"),
            ("Research update", "A small pilot comparing three preprocessing strategies", "UPDATE"),
            ("Methods / workflow", "From raw streams to a reviewed analysis table", "METHODS"),
            ("Results + takeaway", "The most useful result belongs in one sentence", "RESULTS"),
            ("Comparison", "Two analysis choices with different trade-offs", "DECISION"),
            ("Timeline / next steps", "What happens before the next lab meeting", "PLAN"),
            ("Discussion / decision", "Choose the next validation slice", "DISCUSSION"),
            ("Closing", "Questions, notes, and follow-up", "DISCUSSION"),
        ]
    else:
        specs = [
            ("Meeting topic goes here", "Presenter name  •  Lab / date", "LAB MEETING"),
            ("A focused section of the discussion", "Replace with a short section thesis", "SECTION"),
            ("Research update", "Editable placeholder content for a lab meeting", "UPDATE"),
            ("Methods / workflow", "Editable placeholder content for a lab meeting", "METHODS"),
            ("Results + takeaway", "Editable placeholder content for a lab meeting", "RESULTS"),
            ("Comparison", "Editable placeholder content for a lab meeting", "DECISION"),
            ("Timeline / next steps", "Editable placeholder content for a lab meeting", "PLAN"),
            ("Discussion / decision", "Editable placeholder content for a lab meeting", "DISCUSSION"),
            ("Thank you", "Questions, notes, and follow-up", "DISCUSSION"),
        ]

    for page, (title, subtitle, section) in enumerate(specs, 1):
        slide = library.add_slide(presentation, title, subtitle, section, page)
        if page == 1:
            library.text(slide, title, 0.55, 2.28, 8.8, 0.7, size=34, color="navy", bold=True)
            library.text(slide, subtitle, 0.58, 3.1, 8.0, 0.3, size=17, color="muted_text")
            library.callout(slide, "START HERE", "Use one sentence to state the meeting's purpose", 0.55, 4.0, 5.3, 1.0)
        elif page == 2:
            library.text(slide, title, 0.55, 2.65, 10.5, 0.7, size=28, color="navy", bold=True)
            library.text(slide, subtitle, 0.58, 3.48, 7.3, 0.35, size=17, color="muted_text")
        elif page == 3:
            columns = library.three_columns(slide)
            for heading, x in zip(("Question", "Progress", "Open issue"), columns):
                library.text(slide, heading, x + 0.22, 2.58, 3.2, 0.35, size=16, color="navy", bold=True)
            if demo:
                library.bullets(slide, ["Which stream is reliable enough to analyze?", "What changed since last review?"], 0.75, 3.15, 3.15, 1.6)
                library.bullets(slide, ["Pilot preprocessing is complete", "Review set is ready"], 4.74, 3.15, 3.15, 1.6)
                library.bullets(slide, ["Define the validation slice", "Agree on one comparison"], 8.73, 3.15, 3.15, 1.6)
        elif page == 4:
            library.two_columns(slide)
            library.text(slide, "Add a short method description and assumptions", 0.82, 2.58, 5.2, 0.6, size=16, color="navy", bold=True)
            library.bullets(slide, ["Inputs and inclusion rules", "Transformation or model step", "Review and quality checks"], 0.82, 3.35, 5.0, 1.8)
            library.timeline(slide, [("01", "Inputs"), ("02", "Method"), ("03", "Review")], 7.2, 3.4, 4.9)
        elif page == 5:
            library.figure_placeholder(slide, 0.52, 2.05, 7.35, 4.2)
            library.callout(slide, "TAKEAWAY", "State the result in one sentence", 8.2, 2.05, 4.6, 1.55)
            library.metric(slide, "--", "Metric label", 8.2, 4.0)
            library.metric(slide, "--", "Metric label", 10.35, 4.0)
        elif page == 6:
            library.two_columns(slide)
            library.text(slide, "Option A", 0.86, 2.58, 4, 0.3, size=18, color="navy", bold=True)
            library.text(slide, "Option B", 7.13, 2.58, 4, 0.3, size=18, color="navy", bold=True)
            library.bullets(slide, ["Criterion or observation", "Strength or limitation", "What it would require"], 0.86, 3.15, 4.9, 2.0)
            library.bullets(slide, ["Criterion or observation", "Strength or limitation", "What it would require"], 7.13, 3.15, 4.9, 2.0)
        elif page == 7:
            library.timeline(slide, [("Now", "Current state"), ("Next", "Planned work"), ("Later", "Decision point")])
            library.callout(slide, "NEXT STEP", "Name the next action, owner, and review date", 4.25, 5.05, 4.85, 1.0)
        elif page == 8:
            library.callout(slide, "DECISION NEEDED", "Write the question for the room", 0.52, 2.1, 7.0, 1.45)
            library.card(slide, 7.85, 2.1, 4.95, 3.8, "OPTIONS", "teal")
            library.bullets(slide, ["Option or trade-off", "Option or trade-off", "Recommendation or owner"], 8.15, 2.75, 4.2, 2.0)
        else:
            library.text(slide, title, 0.55, 2.55, 8, 0.65, size=34, color="navy", bold=True)
            library.text(slide, subtitle, 0.58, 3.35, 7, 0.3, size=17, color="muted_text")
            library.pill(slide, "DISCUSSION", 0.58, 4.0, 1.55, 0.28, "teal")

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return presentation


def build_template_deck(path: str | Path) -> Any:
    """Build the placeholder-first editable template deck."""
    return _build_deck(path, demo=False)


def build_demo_deck(path: str | Path) -> Any:
    """Build a generic sample deck that demonstrates the same formats."""
    return _build_deck(path, demo=True)


__all__ = ["Library", "add_text", "build_template_deck", "build_demo_deck"]
